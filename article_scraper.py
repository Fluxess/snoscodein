#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import requests
from bs4 import BeautifulSoup
import time
import os
from docx import Document


class ArticleScraper:
    def __init__(self, username, password):
        self.username = username
        self.password = password
        self.session = requests.Session()
        self.base_url = 'https://www.iimes.ru'
        self.login_url = self.base_url + '/wp-login.php'
        self.session.headers.update({
            'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36',
            'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8',
            'Accept-Language': 'ru-RU,ru;q=0.9,en;q=0.8',
        })

    def login(self):
        try:
            self.session.get(self.login_url, timeout=10)
            
            login_data = {
                'log': self.username,
                'pwd': self.password,
                'wp-submit': 'Войти',
                'redirect_to': self.base_url,
                'testcookie': '1'
            }

            response = self.session.post(
                self.login_url,
                data=login_data,
                timeout=15,
                allow_redirects=True
            )
            
            cookies = self.session.cookies.get_dict()
            
            if 'wordpress_logged_in' in str(cookies):
                return True
            if 'wp-login' not in response.url:
                return True
                
            return True

        except Exception as e:
            print('Ошибка сети: ' + str(e))
            return False

    def get_article_text(self, url):
        try:
            response = self.session.get(url, timeout=15)
            if response.status_code != 200:
                return [], ''

            soup = BeautifulSoup(response.text, 'html.parser')
            
            # Получаем заголовок
            title = ''
            for sel in ['h1.entry-title', 'h1.post-title', 'h1', '.entry-title', '.post-title']:
                elem = soup.select_one(sel)
                if elem:
                    title = elem.get_text(strip=True)
                    break

            # Удаляем ненужные элементы
            for tag in soup.find_all(['script', 'style', 'nav', 'footer', 'aside', 'form', 'iframe']):
                tag.decompose()

            # Ищем контент
            content = None
            selectors = [
                '.entry-content',
                '.post-content', 
                '.article-content',
                '.content',
                '.single-content',
                '.post-body',
                'article',
                '.entry',
                'main',
                '#content'
            ]
            
            for selector in selectors:
                content = soup.select_one(selector)
                if content:
                    text_check = content.get_text(strip=True)
                    if len(text_check) > 100:
                        break
                    content = None

            if not content:
                content = soup.find('body')

            if not content:
                return [], title

            # Извлекаем параграфы
            paragraphs = []
            for elem in content.find_all(['p']):
                text = elem.get_text(separator=' ')
                text = ' '.join(text.split())
                if text and len(text) > 30:
                    skip_words = ['cookie', 'войти', 'регистр', 'пароль', 'copyright']
                    if not any(w in text.lower() for w in skip_words):
                        paragraphs.append(text)

            return paragraphs, title

        except Exception as e:
            print('Ошибка: ' + str(e))
            return [], ''

    def get_article_urls(self, max_pages=10):
        urls = []
        
        # Собираем статьи с нескольких страниц
        for page in range(1, max_pages + 1):
            try:
                if page == 1:
                    page_url = self.base_url
                else:
                    page_url = self.base_url + '/page/' + str(page) + '/'
                
                print('Сканирую: ' + page_url)
                response = self.session.get(page_url, timeout=10)
                
                if response.status_code != 200:
                    print('Страница не найдена, останавливаюсь')
                    break
                    
                soup = BeautifulSoup(response.text, 'html.parser')

                found_on_page = 0
                for link in soup.find_all('a', href=True):
                    href = link.get('href')
                    if not href:
                        continue
                    if self.base_url in href or href.startswith('/'):
                        if href.startswith('/'):
                            href = self.base_url + href
                        skip = ['wp-login', 'wp-admin', 'feed', 'category', 'tag', 'author', '#', '.jpg', '.png', '.pdf', '/page/']
                        if not any(s in href.lower() for s in skip):
                            if href not in urls and href != self.base_url and href != self.base_url + '/':
                                urls.append(href)
                                found_on_page += 1
                
                print('  Найдено ссылок: ' + str(found_on_page))
                
                if found_on_page == 0:
                    break
                    
                time.sleep(1)

            except Exception as e:
                print('Ошибка: ' + str(e))
                break

        print('Всего уникальных ссылок: ' + str(len(urls)))
        return urls

    def save_to_word(self, paragraphs, title, filename):
        doc = Document()
        
        if title:
            doc.add_heading(title, 0)
        
        for para in paragraphs:
            doc.add_paragraph(para)
        
        doc.save(filename)

    def scrape_articles(self, max_articles=90):
        if not os.path.exists('articles'):
            os.makedirs('articles')

        # Рассчитываем сколько страниц нужно просканировать
        pages_needed = (max_articles // 10) + 2
        urls = self.get_article_urls(max_pages=pages_needed)

        saved = 0
        for i, url in enumerate(urls, 1):
            if saved >= max_articles:
                break
                
            print('[' + str(i) + '/' + str(len(urls)) + '] ' + url)

            paragraphs, title = self.get_article_text(url)
            
            if paragraphs and len(paragraphs) > 0:
                saved += 1
                
                safe_title = title[:50] if title else 'article'
                for char in ['<', '>', ':', '"', '/', '\\', '|', '?', '*']:
                    safe_title = safe_title.replace(char, '')
                safe_title = safe_title.strip()
                if not safe_title:
                    safe_title = 'article'
                
                filename = 'articles/' + str(saved).zfill(3) + '_' + safe_title + '.docx'
                self.save_to_word(paragraphs, title, filename)
                print('    Сохранено: ' + filename)
            else:
                print('    Пропущено')

            time.sleep(1)
        
        print('\n' + '=' * 40)
        print('Сохранено статей: ' + str(saved))
        print('=' * 40)


if __name__ == '__main__':
    USERNAME = 'Valeev'
    PASSWORD = 'iimes864'

    print('=' * 40)
    print('IIMES Article Scraper')
    print('=' * 40)
    
    scraper = ArticleScraper(USERNAME, PASSWORD)

    print('\nАвторизация...')
    if scraper.login():
        print('OK')
        print('\nПоиск статей...')
        scraper.scrape_articles(max_articles=90)
        print('\nГотово!')
        print('Файлы в папке articles/')
    else:
        print('Ошибка авторизации')
